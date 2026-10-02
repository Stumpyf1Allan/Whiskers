"""Whiskers tests.

Grouped by what would go wrong for a person if the test didn't exist, not by module.
Several are ports of lessons Mittens & Pence paid for; those say so.
"""

from __future__ import annotations

import base64
import datetime as dt
import json
import pathlib
import re
import threading
import urllib.error
import urllib.parse
import urllib.request

import pytest

from folio import config, db
from folio.brokers import freetrade, monzo, symbols, tbills, trading212
from folio.engine import (allocation, export, isa, mathx, narrative, plan, planbuilder, playbook,
                          portfolio, rules, sample, signals, stress)
from folio.market import sources, store

ROOT = pathlib.Path(__file__).resolve().parent.parent
STATIC = ROOT / "folio" / "web" / "static"


# ============================================================================ helpers

def add_platform(name="Freetrade ISA", provider="freetrade", flexible=0, cash=None) -> int:
    with db.tx() as c:
        return c.execute("INSERT INTO platforms(name,provider,wrapper,flexible,cash) "
                         "VALUES(?,?,?,?,?)", (name, provider, "isa", flexible, cash)).lastrowid


def add_position(pid, symbol, exchange="US", qty=10.0, cost=None, broker_value=None,
                 valued_at=None, isin=None, name=None, source="manual") -> int:
    with db.tx() as c:
        iid = symbols.upsert_instrument(c, symbol, exchange, name or symbol, isin, None)
        c.execute("INSERT INTO positions(platform_id,instrument_id,quantity,cost,broker_value,"
                  "valued_at,source,updated_at) VALUES(?,?,?,?,?,?,?,?)",
                  (pid, iid, qty, cost, broker_value, valued_at, source, db.now()))
    return iid


def price_series(key, values, currency="GBP", end=None):
    end = end or dt.date.today()
    pts = [((end - dt.timedelta(days=len(values) - 1 - i)).isoformat(), v)
           for i, v in enumerate(values)]
    store._save(key, {"points": pts, "currency": currency}, "test", replace=True)
    return pts


def sleeve(name, target, **kw) -> int:
    with db.tx() as c:
        return c.execute("INSERT INTO sleeves(name,target,band,ai_share,is_cash) VALUES(?,?,?,?,?)",
                         (name, target, kw.get("band"), kw.get("ai_share"),
                          1 if kw.get("is_cash") else 0)).lastrowid


def days(n, start=dt.date(2024, 1, 1)):
    return [(start + dt.timedelta(days=i)).isoformat() for i in range(n)]


# ============================================================================ maths

def test_moving_average_waits_for_a_full_window():
    assert mathx.sma([1, 2, 3, 4], 2) == [None, 1.5, 2.5, 3.5]


def test_drawdown_is_measured_from_the_trailing_high():
    dd = mathx.drawdowns([100, 120, 90, 120, 60], 252)
    assert dd[2] == pytest.approx(0.25) and dd[3] == 0 and dd[4] == pytest.approx(0.5)


def test_worst_fall_finds_peak_and_trough():
    pts = list(zip(days(6), [100, 150, 75, 140, 160, 120]))
    worst = mathx.max_drawdown(pts)
    assert worst["depth"] == pytest.approx(0.5)
    assert (worst["peak"], worst["trough"]) == (pts[1][0], pts[2][0])


def test_a_change_needs_enough_history_rather_than_guessing():
    pts = list(zip(days(40), range(1, 41)))
    assert mathx.change_over(pts, 365) is None
    assert mathx.change_over(pts, 30) == pytest.approx(40 / 10 - 1)


def test_rates_change_in_points_not_per_cent_of_themselves():
    pts = list(zip(days(400), [5.0] * 35 + [3.0] * 365))
    assert mathx.points_change(pts, 365) == pytest.approx(-2.0)


def test_thinning_keeps_both_ends():
    pts = list(zip(days(1000), range(1000)))
    out = mathx.thin(pts, 100)
    assert len(out) == 100 and out[0] == pts[0] and out[-1] == pts[-1]


# ============================================================================ tickers

@pytest.mark.parametrize("raw,expected", [
    ("PCGHl_EQ", ("PCGH", "LSE")), ("SPXPl_EQ", ("SPXP", "LSE")),
    ("BRK_B_US_EQ", ("BRK-B", "US")), ("GOOGL_US_EQ", ("GOOGL", "US")),
    ("CLA_US_EQ", ("CLA", "US")), ("COPAl_EQ", ("COPA", "LSE")),
])
def test_trading212_tickers_are_decoded(raw, expected):
    assert symbols.parse_t212(raw) == expected


@pytest.mark.parametrize("symbol,exchange,expected", [
    ("VWRP", "LSE", "VWRP.L"), ("SN.", "LSE", "SN.L"), ("AV.", "LSE", "AV.L"),
    ("BT.A", "LSE", "BT-A.L"), ("BRK.B", "US", "BRK-B"), ("NVDA", "US", "NVDA"),
])
def test_price_symbols_follow_yahoo_conventions(symbol, exchange, expected):
    assert symbols.market_symbol(symbol, exchange) == expected


def test_one_spelling_for_matching_companies():
    assert symbols.alias_key("brk.b") == symbols.alias_key("BRK-B") == "BRK-B"
    assert symbols.alias_key("SN.") == "SN"


def test_the_same_isin_on_two_platforms_is_one_instrument():
    a, b = add_platform("A"), add_platform("B", "trading212")
    i1 = add_position(a, "GOOGL", isin="US02079K3059")
    i2 = add_position(b, "GOOGL", isin="US02079K3059")
    assert i1 == i2


# ============================================================================ prices

def _chart(ccy, closes, live):
    return {"chart": {"result": [{"meta": {"currency": ccy, "gmtoffset": 3600,
                                           "regularMarketPrice": live,
                                           "regularMarketTime": 1758800000},
                                  "timestamp": [1758700000 + 86400 * i for i in range(len(closes))],
                                  "indicators": {"quote": [{"close": closes}]}}]}}


def test_pence_are_turned_into_pounds_case_sensitively():
    got = sources.parse_yahoo_chart(_chart("GBp", [1723.5, None, 1800.0], 1810.0))
    assert got["currency"] == "GBP"
    assert [v for _, v in got["points"]] == [17.235, 18.0] and got["live"] == 18.1
    pounds = sources.parse_yahoo_chart(_chart("GBP", [17.2], 17.3))
    assert pounds["points"][0][1] == 17.2        # "GBP" is pounds already


def test_fred_csv_handles_both_headers_and_gaps():
    old = "DATE,VIXCLS\n2026-01-02,14.5\n2026-01-05,.\n2026-01-06,15.1\n"
    new = "observation_date,VIXCLS\n2026-01-02,14.5\n2026-01-05,\n"
    assert sources.parse_fred_csv(old) == [("2026-01-02", 14.5), ("2026-01-06", 15.1)]
    assert sources.parse_fred_csv(new) == [("2026-01-02", 14.5)]


def test_no_exchange_rate_means_no_pound_value_not_a_dollar_one():
    assert store.to_gbp(100, "USD") is None


# ============================================================================ Freetrade

FT_OLD = ("Title,Type,Timestamp,Account Currency,Total Amount,Buy / Sell,Ticker,ISIN,"
          "Price per Share in Account Currency,Stamp Duty,Quantity,Venue,Order ID,"
          "Instrument Currency,FX Fee Amount\n")
FT_NEW = FT_OLD.replace("Total Amount,", "Total Amount in Account Currency,")


def _ft(header, rows):
    return (header + "".join(",".join(r) + "\n" for r in rows)).encode()


ROWS = [
    ["Top up", "TOP_UP", "2025-04-05T23:30:00.000Z", "GBP", "1000", "", "", "", "", "", "", "", "", "", ""],
    ["Vanguard", "ORDER", "2025-04-07T09:00:00.000Z", "GBP", "600", "BUY", "VWRP", "IE00BK5BQT80",
     "100", "0", "6", "London Stock Exchange", "o1", "GBP", "0"],
    ["Vanguard", "ORDER", "2025-05-07T09:00:00.000Z", "GBP", "300", "SELL", "VWRP", "IE00BK5BQT80",
     "150", "0", "2", "London Stock Exchange", "o2", "GBP", "0"],
    ["Free share", "FREESHARE_ORDER", "2025-06-01T09:00:00.000Z", "GBP", "25", "BUY", "SN.",
     "GB0009223206", "12.5", "0", "2", "London Stock Exchange", "o3", "GBP", "0"],
]


def test_utc_timestamps_are_dated_in_uk_time():
    assert freetrade.uk_date("2025-04-05T23:30:00.000Z") == "2025-04-06"   # BST: next tax year
    assert freetrade.uk_date("2025-01-10T23:30:00Z") == "2025-01-10"       # GMT: same day
    assert freetrade.uk_date("2025-10-26T00:30:00Z") == "2025-10-26"       # BST until 01:00


@pytest.mark.parametrize("header", [FT_OLD, FT_NEW])
def test_both_spellings_of_the_amount_column_import_with_amounts(header):
    pid = add_platform()
    rep = freetrade.import_csv(pid, _ft(header, ROWS))
    assert rep["trades"] == 3 and rep["cash"] == 1
    pos = {r["symbol"]: r for r in db.rows(
        "SELECT i.symbol, p.quantity, p.cost FROM positions p JOIN instruments i ON i.id=p.instrument_id")}
    assert pos["VWRP"]["quantity"] == 4
    assert pos["VWRP"]["cost"] == pytest.approx(400)      # pooled: a third of 600 went with the sale
    assert pos["SN."]["cost"] == 0                         # a free share cost nothing
    assert isa.allowance(2025)["used"] == 1000             # 23:30 UTC on 5 April is 6 April in the UK


def test_importing_the_same_file_twice_changes_nothing():
    pid = add_platform()
    freetrade.import_csv(pid, _ft(FT_NEW, ROWS))
    again = freetrade.import_csv(pid, _ft(FT_NEW, ROWS))
    assert again["trades"] == 0 and again["cash"] == 0
    assert db.scalar("SELECT COUNT(*) FROM trades") == 3


def test_a_hand_correction_survives_the_next_import():
    pid = add_platform()
    freetrade.import_csv(pid, _ft(FT_NEW, ROWS))
    iid = db.scalar("SELECT id FROM instruments WHERE symbol='VWRP'")
    freetrade.adjust_quantity(pid, iid, 5.5)
    freetrade.import_csv(pid, _ft(FT_NEW, ROWS))
    assert db.scalar("SELECT quantity FROM positions WHERE instrument_id=?", (iid,)) == 5.5


def test_a_file_from_somewhere_else_is_refused_plainly():
    with pytest.raises(freetrade.NotFreetrade):
        freetrade.read_rows(b"Date,Description,Amount\n2025-01-01,Coffee,-3\n")


# ============================================================================ Trading 212

POSITION = {"instrument": {"ticker": "VUAGl_EQ", "name": "Vanguard S&P 500", "isin": "IE00BFMXXD54",
                           "currency": "GBX"},
            "quantity": 12.5, "averagePricePaid": 8800, "currentPrice": 9950,
            "walletImpact": {"currency": "GBP", "currentValue": 1243.75, "totalCost": 1100,
                             "unrealizedProfitLoss": 143.75, "fxImpact": 0}}


def test_positions_are_read_in_pounds_from_wallet_impact():
    p = trading212.parse_position(POSITION)
    assert (p["symbol"], p["exchange"], p["value_gbp"], p["cost_gbp"]) == ("VUAG", "LSE", 1243.75, 1100)


def test_orders_are_read_from_the_current_nested_shape():
    item = {"order": {"id": 7, "side": "BUY", "status": "FILLED", "createdAt": "2026-03-01T10:00:00Z",
                      "instrument": {"ticker": "NVDA_US_EQ", "name": "NVIDIA", "isin": "US67066G1040",
                                     "currency": "USD"}},
            "fill": {"id": 9, "filledAt": "2026-03-01T10:00:02Z", "price": 120.5, "quantity": 2,
                     "type": "TRADE", "walletImpact": {"currency": "GBP", "netValue": -190.2,
                                                       "fxRate": 0.79, "taxes": [{"quantity": 0.29}]}}}
    o = trading212.parse_order(item)
    assert (o["side"], o["quantity"], o["value_gbp"], o["fees_gbp"], o["ref"]) == \
        ("BUY", 2, 190.2, 0.29, "7:9")
    assert trading212.parse_order({"order": {"id": 1, "side": "BUY"}, "fill": None}) is None
    item["fill"]["type"] = "FOP"
    assert trading212.parse_order(item)["side"] == "TRANSFER_IN"


def test_the_old_flat_order_shape_still_reads():
    o = trading212.parse_order({"ticker": "AAPL_US_EQ", "filledQuantity": -3, "status": "FILLED",
                                "fillPrice": 200, "fillCost": 470, "dateExecuted": "2024-02-02T12:00:00Z",
                                "id": 5})
    assert (o["side"], o["quantity"], o["date"]) == ("SELL", 3, "2024-02-02")


def test_withdrawals_are_negative_and_interest_positive():
    assert trading212.parse_transaction({"type": "WITHDRAW", "amount": 50, "dateTime": "2026-01-01T00:00:00Z"})["amount"] == -50
    assert trading212.parse_transaction({"type": "INTEREST_ON_FREE_CASH", "amount": -1.2, "dateTime": "2026-01-01"})["amount"] == 1.2


def test_auth_header_is_basic_key_colon_secret_on_one_line():
    h = trading212.Client("KEY", "SECRET")._headers()["Authorization"]
    assert h == "Basic " + base64.b64encode(b"KEY:SECRET").decode() and "\n" not in h


def test_401_and_403_get_opposite_advice():
    """Mittens & Pence once read a 401 as a missing permission and told Allan to grant
    'Orders - Execute'. A 401 is the key; a 403 is a permission — and never that one."""
    assert "not permissions" in trading212.explain(401)
    msg = trading212.explain(403)
    assert "NOT need 'Orders - Execute'" in msg


def test_the_trading212_client_cannot_place_an_order():
    src = (ROOT / "folio" / "brokers" / "trading212.py").read_text(encoding="utf-8")
    assert "net.post" not in src and "method=" not in src
    assert not re.search(r'"/api/v0/equity/orders', src)
    net_src = (ROOT / "folio" / "net.py").read_text(encoding="utf-8")
    assert "def post" not in net_src


class FakeT212:
    def __init__(self, summary):
        self.summary = summary

    def get(self, path):
        return self.summary if path.endswith("summary") else [POSITION]


def test_sync_stores_positions_and_cash(monkeypatch):
    pid = add_platform("T212 ISA", "trading212")
    monkeypatch.setattr(trading212, "_client", lambda _: FakeT212(
        {"totalValue": 1500, "cash": {"availableToTrade": 200, "inPies": 50, "reservedForOrders": 6.25}}))
    monkeypatch.setattr(trading212.time, "sleep", lambda s: None)
    rep = trading212.sync(pid, history=False)
    assert rep["positions"] == 1
    assert db.scalar("SELECT cash FROM platforms WHERE id=?", (pid,)) == 256.25
    assert db.scalar("SELECT broker_value FROM positions") == 1243.75


def test_cash_in_an_unexpected_shape_stays_unknown(monkeypatch):
    """Mittens & Pence read Trading 212's cash under the wrong names and so always
    showed it as unknown; worse would be showing it as £0. Neither happens here."""
    pid = add_platform("T212 ISA", "trading212")
    monkeypatch.setattr(trading212, "_client", lambda _: FakeT212({"free": 12}))
    monkeypatch.setattr(trading212.time, "sleep", lambda s: None)
    rep = trading212.sync(pid, history=False)
    assert db.scalar("SELECT cash FROM platforms WHERE id=?", (pid,)) is None
    assert rep["warnings"]


# ============================================================================ valuation

def test_a_holding_with_no_price_is_unknown_not_zero():
    pid = add_platform(cash=100)
    add_position(pid, "MYSTERY", qty=5)
    t = portfolio.totals()
    assert t["unpriced"] == 1 and t["invested"] is None and t["total"] == 100


def test_no_cash_figure_is_not_zero_cash():
    pid = add_platform()
    add_position(pid, "VWRP", "LSE", broker_value=1000, valued_at=db.now())
    assert portfolio.totals()["cash"] is None


def test_market_prices_are_converted_at_the_days_rate():
    pid = add_platform(cash=0)
    add_position(pid, "NVDA", qty=10, cost=900)
    price_series("SEC:NVDA", [100.0] * 30, "USD")
    price_series("FX:USD", [1.25] * 30)
    x = portfolio.positions()[0]
    assert x["value"] == pytest.approx(800) and x["gain"] == pytest.approx(-100)


def test_same_isin_elsewhere_prices_a_holding_with_no_market_price():
    a, b = add_platform("FT"), add_platform("T212", "trading212")
    add_position(a, "XPO", isin="US9837931008", qty=2, name="XPO")
    add_position(b, "XPO", isin="US9837931008", qty=4, broker_value=400, valued_at=db.now())
    rows = {r["platform"]: r for r in portfolio.positions()}
    assert rows["FT"]["value"] == pytest.approx(200) and rows["FT"]["value_source"] == "broker-same-isin"


def test_backdated_history_ends_at_todays_value_and_reports_coverage():
    pid = add_platform(cash=0)
    add_position(pid, "AAA", "LSE", qty=10, name="Covered")
    add_position(pid, "BBB", "LSE", qty=1, broker_value=100, valued_at=db.now(), name="No history")
    price_series("SEC:AAA.L", [10.0 + i for i in range(60)])
    b = portfolio.backdated(365)
    assert b["points"][-1][1] == pytest.approx(690)
    assert b["missing"] == ["No history"] and b["coverage"] == pytest.approx(690 / 790)


# ============================================================================ allocation

@pytest.mark.parametrize("target,band", [(40, 5), (23, 5), (8, 2), (2, 0.5), (0, 0.25)])
def test_bands_follow_the_5_25_rule(target, band):
    assert allocation.band_for(target) == band


def _view(actuals: dict[str, tuple[float, float]], total=1000.0):
    rows = [{"id": i, "name": n, "target": t, "value": a / 100 * total, "actual": a,
             "colour": "#000"} for i, (n, (t, a)) in enumerate(actuals.items())]
    return {"sleeves": rows, "extra": [], "total": total}


def test_the_hundred_squares_always_add_up_to_a_hundred():
    v = _view({"A": (33.3, 33.3), "B": (33.3, 33.3), "C": (33.4, 33.4)})
    assert len(allocation.waffle(v)) == 100
    v = _view({"A": (50, 12.5), "B": (50, 87.5)})
    assert len(allocation.waffle(v)) == 100


def test_new_money_goes_to_the_underweight_sleeves_first():
    v = _view({"Over": (50, 60), "Under": (50, 40)})
    split = {s["name"]: s["amount"] for s in allocation.new_money(100, v)}
    assert split["Under"] == pytest.approx(100) and split["Over"] == 0
    split = {s["name"]: s["amount"] for s in allocation.new_money(1000, v)}
    assert sum(split.values()) == pytest.approx(1000) and min(split.values()) >= 0


def test_status_goes_amber_past_half_the_band_and_red_past_the_band():
    pid = add_platform(cash=0)
    s = sleeve("Core", 40)
    for sym, value in (("A", 425), ("B", 575)):
        iid = add_position(pid, sym, broker_value=value, valued_at=db.now())
        db.execute("UPDATE instruments SET sleeve_id=? WHERE id=?", (s if sym == "A" else None, iid))
    view = allocation.view(portfolio.positions(), portfolio.totals())
    core = view["sleeves"][0]
    assert core["band"] == 5 and core["status"] == "ok"             # 42.5%: within half the band
    assert view["extra"][0]["name"] == "Not in a sleeve yet"          # never silently dropped


# ============================================================================ rules

def _alphabet_at(pct):
    pid, other = add_platform("FT", cash=0), add_platform("T212", "trading212", cash=0)
    add_position(pid, "GOOGL", broker_value=pct * 5, valued_at=db.now(), isin="US02079K3059")
    add_position(other, "GOOGL", broker_value=pct * 5, valued_at=db.now(), isin="US02079K3059")
    add_position(other, "VWRP", "LSE", broker_value=1000 - pct * 10, valued_at=db.now())
    db.execute("INSERT INTO rules(kind,subject,trigger_pct,target_pct,basis) "
               "VALUES('company_cap','Alphabet',9.3,8,'direct')")
    pos = portfolio.positions()
    return rules.rule_status(pos, portfolio.totals(pos))[0]


def test_alphabet_ratchet_counts_both_platforms_and_sizes_the_trim():
    s = _alphabet_at(9.5)
    assert s["now"] == pytest.approx(9.5) and s["level"] == "red"
    assert s["trim"] == pytest.approx(15)             # 1.5% of £1,000 back to 8%


@pytest.mark.parametrize("pct,level", [(8.0, "green"), (8.7, "amber"), (9.31, "red")])
def test_alphabet_ratchet_levels(pct, level):
    assert _alphabet_at(pct)["level"] == level


def test_a_platform_over_the_fscs_limit_is_red_and_near_it_amber():
    big, near = add_platform("Big", cash=0), add_platform("Near", cash=0)
    add_position(big, "A", broker_value=86000, valued_at=db.now())
    add_position(near, "B", broker_value=77000, valued_at=db.now())
    pos = portfolio.positions()
    tot = portfolio.totals(pos)
    warns = {w["id"]: w["level"] for w in rules.evaluate(pos, tot, allocation.view(pos, tot))}
    assert warns[f"fscs:{big}"] == "red" and warns[f"fscs:{near}"] == "amber"


# ============================================================================ market lights

def _levels(defn, values):
    return signals.levels_for(defn, list(zip(days(len(values)), values)))


def test_trend_light_goes_red_only_once_the_average_itself_falls():
    trend = signals.DEFS[0]
    rising = _levels(trend, [100 + i for i in range(260)])
    assert rising[-1][1] == "green"
    dipped = _levels(trend, [100 + i for i in range(260)] + [300] * 5 + [150] * 5)
    assert dipped[-1][1] == "amber"                   # below the average, average still rising
    crashed = _levels(trend, [100 + i for i in range(260)] + [100] * 120)
    assert crashed[-1][1] == "red"


def test_fall_from_high_uses_its_own_thresholds():
    sox = next(d for d in signals.DEFS if d["key"] == "sox_fall")
    assert _levels(sox, [100] * 100 + [85] * 5)[-1][1] == "amber"
    assert _levels(sox, [100] * 100 + [79] * 5)[-1][1] == "red"


@pytest.mark.parametrize("ai,mkt,regime", [
    (["green"] * 4, ["green"] * 2, "calm"), (["amber", "green", "green", "green"], ["green"] * 2, "watch"),
    (["red", "red", "green", "green"], ["green"] * 2, "turning"),
    (["green"] * 4, ["red", "red", "green", "green"], "turning"),      # markets can turn without AI
    (["red", "red", "amber", "green"], ["red", "red", "green", "green"], "stress"), ([], [], "none"),
])
def test_regimes(ai, mkt, regime):
    assert signals.regime_of(ai, mkt)[0] == regime


def test_lights_with_no_data_are_grey_never_green():
    got = signals.evaluate()
    assert got["regime"] == "none" and {x["level"] for x in got["lights"]} == {"none"}


# ============================================================================ ISA

def test_the_tax_year_turns_on_the_sixth_of_april():
    assert isa.tax_year(dt.date(2026, 4, 5)) == 2025 and isa.tax_year(dt.date(2026, 4, 6)) == 2026
    assert isa.label(2025) == "2025/26"


def test_flexible_isas_net_off_withdrawals_and_transfers_are_not_subscriptions():
    flex, fixed = add_platform("T212", "trading212", flexible=1), add_platform("FT")
    with db.tx() as c:
        for pid, kind, amt in ((flex, "DEPOSIT", 5000), (flex, "WITHDRAWAL", -2000),
                               (fixed, "DEPOSIT", 3000), (fixed, "WITHDRAWAL", -1000),
                               (fixed, "TRANSFER", 9000)):
            c.execute("INSERT INTO cash_moves(platform_id,happened_on,kind,amount_gbp,fp) "
                      "VALUES(?,?,?,?,?)", (pid, "2026-05-01", kind, amt, f"{pid}{kind}"))
    a = isa.allowance(2026)
    assert a["used"] == 3000 + 3000 and a["transfers"] == 9000


# ============================================================================ plans and sample

def test_a_plan_round_trips():
    plan.import_plan({"whiskers_plan": 1,
                      "sleeves": [{"name": "Core", "target": 60}, {"name": "Cash", "target": 40, "is_cash": True}],
                      "rules": [{"kind": "company_cap", "subject": "Alphabet", "trigger": 9.3, "target": 8}],
                      "companies": {"GOOGL": "Alphabet"}, "themes": {"XPO": "Freight"}})
    out = plan.export_plan()
    assert [s["name"] for s in out["sleeves"]] == ["Core", "Cash"]
    assert out["rules"][0]["trigger"] == 9.3
    pid = add_platform()
    add_position(pid, "XPO")                         # arriving after the plan still gets its theme
    assert db.scalar("SELECT theme FROM instruments WHERE symbol='XPO'") == "Freight"


@pytest.mark.parametrize("bad", [
    {"whiskers_plan": 1, "sleeves": [{"name": "X", "target": 140}]},
    {"whiskers_plan": 99}, {"sleeves": []}, {"whiskers_plan": 1, "rules": [{"kind": "buy_everything"}]},
])
def test_bad_plans_are_refused_with_a_reason(bad):
    with pytest.raises(plan.PlanError):
        plan.import_plan(bad)


def test_loading_the_sample_twice_does_not_build_a_second_portfolio():
    """Mittens & Pence's sample doubled the household when loaded twice."""
    sample.load()
    sample.load()
    assert db.scalar("SELECT COUNT(*) FROM platforms") == 2
    sample.remove()
    assert db.scalar("SELECT COUNT(*) FROM platforms") == 0
    assert db.scalar("SELECT COUNT(*) FROM instruments") == 0


def test_stress_test_is_arithmetic_on_todays_holdings():
    exp = {"ai_value": 400.0, "unknown_value": 0.0}
    out = stress.scenario(exp, {"total": 1000.0, "cash": 100.0}, 0.5, 0.1)
    assert out["loss"] == pytest.approx(200 + 50) and out["after"] == pytest.approx(750)


# ============================================================================ the local server

@pytest.fixture()
def server():
    from folio.web import server as srv
    httpd, port = srv.serve(port=0 or srv._free_port(18766), block=False)
    yield port
    httpd.shutdown()


def _req(port, path, method="GET", headers=None, body=None):
    req = urllib.request.Request(f"http://127.0.0.1:{port}{path}", method=method,
                                 data=body, headers=headers or {})
    try:
        with urllib.request.urlopen(req, timeout=5) as r:
            return r.status, r.read(), dict(r.headers)
    except urllib.error.HTTPError as e:
        return e.code, e.read(), dict(e.headers)


def test_other_websites_cannot_use_the_local_server(server):
    ok = {"Content-Type": "application/json"}
    assert _req(server, "/api/sample/load", "POST", ok, b"{}")[0] == 200          # no Origin: a tool
    for bad in ({"Origin": "https://evil.example"}, {"Origin": "null"},
                {"Origin": "http://127.0.0.1:1"},          # another local app, e.g. Mittens & Pence
                {"Sec-Fetch-Site": "cross-site"}):
        assert _req(server, "/api/sample/remove", "POST", {**ok, **bad}, b"{}")[0] == 403, bad
    assert _req(server, "/api/overview", headers={"Host": "rebound.example"})[0] == 403
    own = {**ok, "Origin": f"http://127.0.0.1:{server}"}
    assert _req(server, "/api/sample/remove", "POST", own, b"{}")[0] == 200


def test_static_files_cannot_be_escaped(server):
    status, body, _ = _req(server, "/../server.py")
    assert b"def serve" not in body
    status, body, _ = _req(server, "/%2e%2e/server.py")
    assert b"def serve" not in body


def test_pages_carry_a_content_security_policy(server):
    status, body, headers = _req(server, "/")
    assert status == 200 and "default-src 'self'" in headers.get("Content-Security-Policy", "")


def test_every_screen_answers_on_an_empty_copy(server):
    for path in ("/api/bootstrap", "/api/overview", "/api/holdings", "/api/allocation", "/api/ai",
                 "/api/markets", "/api/activity", "/api/settings", "/api/update", "/api/refresh"):
        status, body, _ = _req(server, path)
        assert status == 200, (path, body[:200])
        assert "error" not in json.loads(body), path


def test_every_screen_answers_with_the_sample_loaded(server):
    sample.load()
    for path in ("/api/overview", "/api/holdings", "/api/allocation", "/api/ai", "/api/markets",
                 "/api/activity", "/api/settings", "/api/holding/1"):
        status, body, _ = _req(server, path)
        assert status == 200, (path, body[:300])


# ============================================================================ front-end contracts

def _static(name):
    return (STATIC / name).read_text(encoding="utf-8")


def test_the_interface_loads_nothing_from_the_internet():
    for name in ("index.html", "app.js", "charts.js", "styles.css"):
        src = _static(name)
        assert not re.search(r'<script[^>]+src=["\']https?:', src), name
        assert not re.search(r'<link[^>]+href=["\']https?:', src), name
        assert "@import" not in src and not re.search(r"url\(\s*['\"]?https?:", src), name
        assert not re.search(r"fetch\(\s*['\"`]https?:", src), name


def test_hide_amounts_is_one_class_on_the_page():
    css = _static("styles.css")
    assert re.search(r"html\.hide\s+\.pv\s*\{[^}]*blur", css)
    app = _static("app.js")
    assert "class=\"pv\"" in app or "class='pv'" in app or "class=\"pv " in app


def test_no_chart_uses_an_svg_title_which_nothing_can_blur():
    for name in ("app.js", "charts.js"):
        code = re.sub(r"/\*.*?\*/", "", _static(name), flags=re.S)
        code = re.sub(r"(?m)^\s*//.*$", "", code)
        assert "<title" not in code, name


def test_every_button_action_has_a_handler():
    app = _static("app.js")
    used = set(re.findall(r'data-act="([a-z0-9-]+)"', app)) | set(re.findall(r"data-act=\\?'([a-z0-9-]+)", app))
    defined = set(re.findall(r"ACTS\['([a-z0-9-]+)'\]\s*=", app)) | set(re.findall(r"ACTS\.([a-zA-Z0-9_]+)\s*=", app))
    defined |= set(re.findall(r"^\s*'([a-z0-9-]+)':", app, re.M))
    missing = sorted(used - defined)
    assert not missing, f"buttons with no handler: {missing}"


def test_the_api_never_uses_error_for_anything_but_failure():
    """A top-level "error" makes the front end treat a reply as a failed request —
    Mittens & Pence lost its update explanations that way."""
    src = (ROOT / "folio" / "web" / "server.py").read_text(encoding="utf-8")
    routes = src.split("class Handler")[0]
    assert '"error":' not in routes


# ============================================================================ packaging

def test_text_is_always_read_and_written_with_an_encoding():
    """Windows reads text as cp1252 unless told otherwise: 18 bare read_text() calls
    stopped a Mittens & Pence build."""
    bad = []
    for f in list((ROOT / "folio").rglob("*.py")) + list((ROOT / "tools").glob("*.py")):
        for n, line in enumerate(f.read_text(encoding="utf-8").splitlines(), 1):
            if re.search(r"\.read_text\(\s*\)", line) or re.search(r"\.write_text\([^)]*\)", line) \
                    and "encoding" not in line:
                bad.append(f"{f.name}:{n}")
            if re.search(r"\bopen\([^)]*['\"][rwa]t?['\"]", line) and "encoding" not in line and "b'" not in line \
                    and '"rb"' not in line and "'rb'" not in line and '"wb"' not in line:
                bad.append(f"{f.name}:{n}")
    assert not bad, bad


def test_the_spec_points_at_real_files_and_never_excludes_an_aliased_module():
    spec = (ROOT / "build" / "whiskers.spec").read_text(encoding="utf-8")
    assert 'PKG = ROOT / "folio"' in spec and "run_whiskers.py" in spec
    aliased = re.search(r"ALIASED_BY_PYINSTALLER = \(([^)]*)\)", spec).group(1)
    excludes = re.search(r"excludes = \[(.*?)\]", spec, re.S).group(1)
    for name in re.findall(r'"([^"]+)"', aliased):
        assert f'"{name}"' not in excludes
    assert (ROOT / "run_whiskers.py").exists()
    for icon in ("whiskers.ico", "whiskers.icns", "whiskers.png"):
        assert (ROOT / "folio" / "resources" / icon).exists(), icon


def test_windows_scripts_have_windows_line_endings():
    for f in (ROOT / "build" / "build_windows.bat", ROOT / "Start Whiskers.bat"):
        raw = f.read_bytes()
        assert raw.count(b"\n") == raw.count(b"\r\n"), f.name
    rules = (ROOT / ".gitattributes").read_text(encoding="utf-8")
    assert "*.bat  text eol=crlf" in rules and "*.sh       text eol=lf" in rules


def test_the_version_is_the_same_everywhere_it_is_read():
    assert re.fullmatch(r"\d+\.\d+\.\d+", config.APP_VERSION)
    assert config.UPDATE_MANIFEST_URL.startswith("https://github.com/Stumpyf1Allan/whiskers/")


def test_the_release_workflow_matches_the_project():
    yaml = pytest.importorskip("yaml")
    text = (ROOT / ".github" / "workflows" / "release.yml").read_text(encoding="utf-8")
    yaml.safe_load(text)
    assert "build/whiskers.spec" in text and "WHISKERS_ONEFILE" in text and "pytest" in text
    assert not re.search(r"(?i)kestrel|mittens", text)
    assert (ROOT / "build" / "Read me first.txt").exists()


def test_plan_settings_wait_for_a_holding_that_arrives_later():
    """Trading 212 still lists Ouster as CLA, its old SPAC ticker; the plan says its
    price lives under OUST. The plan is loaded first and the holding turns up later."""
    plan.import_plan({"whiskers_plan": 1, "sleeves": [{"name": "Speculative", "target": 2}],
                      "instruments": {"CLA": {"market_symbol": "OUST", "sleeve": "Speculative"},
                                      "brk.b": {"ai_share": 0}}})
    with db.tx() as c:
        iid = symbols.upsert_instrument(c, *symbols.parse_t212("CLA_US_EQ"), "Ouster", None, "USD")
        brk = symbols.upsert_instrument(c, *symbols.parse_t212("BRK_B_US_EQ"), "Berkshire", None, "USD")
    row = db.one("SELECT i.market_symbol, s.name AS sleeve FROM instruments i "
                 "LEFT JOIN sleeves s ON s.id=i.sleeve_id WHERE i.id=?", (iid,))
    assert (row["market_symbol"], row["sleeve"]) == ("OUST", "Speculative")
    assert db.scalar("SELECT ai_share FROM instruments WHERE id=?", (brk,)) == 0


# ============================================================================ Treasury bills

@pytest.mark.parametrize("name,isin,expected", [
    ("UK Treasury Bill", "GB00BSGHFL62", True), ("UK T-Bill 30/09/2024", "GB00BSGHFL62", True),
    ("UKTB 0 09/30/24", None, True), ("4 1/8% Treasury Gilt 2027", "GB00BSQNRD19", False),
    ("iShares $ Treasury Bond 1-3yr", "IE00B14X4S71", False), ("Vanguard FTSE All-World", None, False),
])
def test_treasury_bills_are_recognised_narrowly(name, isin, expected):
    """A gilt mistaken for a bill would be dropped after five weeks, so the net is tight."""
    assert tbills.is_tbill(name, isin, isin) is expected


@pytest.mark.parametrize("name,bought,expected", [
    ("UK Treasury Bill 30/09/2024", "2024-09-02", dt.date(2024, 9, 30)),
    ("UKTB 0 09/30/24", "2024-09-02", dt.date(2024, 9, 30)),          # Bloomberg's month-first
    ("Treasury Bill 30 Sep 2024", "2024-09-02", dt.date(2024, 9, 30)),
    ("Treasury Bill 30sep2024", "2024-09-02", dt.date(2024, 9, 30)),
    ("UK Treasury Bill", "2024-09-02", dt.date(2024, 10, 7)),          # no date: a safe bound
    ("UK Treasury Bill 30/09/2031", "2024-09-02", dt.date(2024, 10, 7)),  # nonsense for a bill
])
def test_a_bills_maturity_is_read_sensibly(name, bought, expected):
    assert tbills.maturity(name, bought) == expected


def test_no_price_feed_is_asked_about_an_isin():
    assert symbols.market_symbol("GB00BSGHFL62", "LSE") == ""


def _bill_row(isin, when, amount, title="UK Treasury Bill"):
    return [title, "ORDER", f"{when}T09:00:00.000Z", "GBP", str(amount), "BUY", isin, isin,
            "0.996", "0", str(round(amount / 0.996, 2)), "", f"b{isin}", "GBP", "0"]


def test_matured_bills_drop_out_and_a_live_one_counts_at_cost_as_cash():
    sleeve("Cash", 4, is_cash=True)
    pid = add_platform(cash=0)
    live_day = (dt.date.today() - dt.timedelta(days=5)).isoformat()
    rows = [_bill_row("GB00BSGHFL62", "2024-09-02", 1000), _bill_row("GB00BSGHQV34", "2024-10-07", 1004),
            _bill_row("GB00BSGQD757", live_day, 1010)]
    freetrade.import_csv(pid, _ft(FT_NEW, rows))
    pos = portfolio.positions()
    assert [x["isin"] for x in pos] == ["GB00BSGQD757"]
    assert (pos[0]["value"], pos[0]["value_source"]) == (1010, "at-cost")
    assert pos[0]["sleeve"] == "Cash"
    # three bought, two repaid: only the live bill's cost is still out of the cash
    assert freetrade.cash_estimate(pid) == pytest.approx(-(1000 + 1004 + 1010) + (1000 + 1004))
    assert portfolio.totals(pos)["unpriced"] == 0


def test_series_nothing_holds_any_more_are_forgotten():
    pid = add_platform()
    iid = add_position(pid, "VWRP", "LSE")
    db.execute("UPDATE instruments SET market_symbol='GB00BSGHFL62.L' WHERE id=?", (iid,))
    for key in ("SEC:GB00BSGHFL62.L", "SEC:CLA", "REF:SP500"):
        store._fail(key, "Yahoo doesn't recognise it")
    store.repair()
    assert db.scalar("SELECT market_symbol FROM instruments WHERE id=?", (iid,)) is None
    store.prune(store.wanted_keys())
    left = {r["key"] for r in db.rows("SELECT key FROM series_meta")}
    assert left == {"REF:SP500"}


def test_the_connection_reports_why_it_is_plain():
    problem = __import__("folio.net", fromlist=["x"]).browser_tls_problem()
    assert problem is None or ":" in problem


# ============================================================================ cash accounts

def test_cash_accounts_sit_beside_the_portfolio_unless_part_of_the_plan():
    isa_ = add_platform("FT", cash=1000)
    add_position(isa_, "VWRP", "LSE", broker_value=9000, valued_at=db.now())
    with db.tx() as c:
        hsbc = c.execute("INSERT INTO platforms(name,provider,wrapper,cash,in_plan) VALUES"
                         "('HSBC','bank','cash',5000,0)").lastrowid
    t = portfolio.totals()
    assert (t["total"], t["cash"], t["outside_cash"], t["everything"]) == (10000, 1000, 5000, 15000)
    db.execute("UPDATE platforms SET in_plan=1 WHERE id=?", (hsbc,))
    t = portfolio.totals()
    assert (t["total"], t["cash"], t["outside_cash"]) == (15000, 6000, None)
    bank = next(p for p in t["platforms"] if p["name"] == "HSBC")
    assert bank["kind"] == "bank" and bank["limit"] == 120000      # deposit protection, not £85k
    assert isa.allowance()["platforms"][0]["name"] == "FT"        # a bank isn't an ISA


# ============================================================================ the playbook

def test_signals_count_and_tiers_follow_the_playbook():
    price_series("REF:US10Y", [5.2] * 30)                   # fires: above 5%
    price_series("REF:VIX", [15.0] * 30)                    # quiet
    playbook.set_manual("hyperscalers", ticked=["Amazon", "Meta", "Nokia"])   # 2 of 4 fires
    playbook.set_manual("mag7", value=31.5)
    got = playbook.evaluate()
    rows = {r["key"]: r for r in got["signals"]}
    assert rows["us10y"]["firing"] and rows["hyperscalers"]["firing"] and not rows["vix"]["firing"]
    assert rows["hyperscalers"]["ticked"] == ["Amazon", "Meta"]
    assert rows["uk10y"]["firing"] is None                  # no data is not "fine"
    assert (got["firing"], got["tier"]["label"]) == (2, "Elevated")


def test_breadth_is_measured_from_the_baseline_date():
    price_series("REF:RSP", [100.0] * 60 + [88.0] * 5)
    price_series("REF:SPY", [100.0] * 65)
    db.set_meta("playbook", {"breadth_baseline": (dt.date.today() - dt.timedelta(days=30)).isoformat()})
    row = next(r for r in playbook.evaluate()["signals"] if r["key"] == "breadth")
    assert row["value"] == pytest.approx(-12.0) and row["firing"]


def test_the_ladder_reads_the_fall_of_todays_holdings():
    pts = list(zip(days(40), [100.0] * 20 + [100 - i for i in range(1, 21)]))   # 20% down at the end
    lad = playbook.ladder({"points": pts}, 3000.0)
    active = next(r for r in lad["rungs"] if r.get("active"))
    assert lad["drop"] == pytest.approx(20) and active["from"] == 20 and lad["third"] == 1000


def test_holding_targets_match_across_spellings_and_say_what_is_missing():
    pid = add_platform(cash=0)
    add_position(pid, "BRK-B", broker_value=250, valued_at=db.now())
    iid = add_position(pid, "CLA", broker_value=100, valued_at=db.now())
    db.execute("UPDATE instruments SET market_symbol='OUST' WHERE id=?", (iid,))
    add_position(pid, "VWRP", "LSE", broker_value=9650, valued_at=db.now())
    plan.import_plan({"whiskers_plan": 1, "targets": {"BRK.B": 2.5, "OUST": 1.35, "AIGC": 1.5}})
    pos = portfolio.positions()
    rows = {r["key"]: r for r in allocation.holding_targets(pos, portfolio.totals(pos)["total"])}
    assert rows["BRK-B"]["status"] == "ok" and rows["BRK-B"]["actual"] == pytest.approx(2.5)
    assert rows["OUST"]["status"] == "out"                   # 1% against 1.35% ± a quarter
    assert rows["AIGC"]["status"] == "not held"


def test_allans_plan_file_loads_whole():
    p = ROOT.parent / "out" / "Allan's plan.json"
    if not p.exists():
        pytest.skip("the personal plan file lives outside the project")
    rep = plan.import_plan(p.read_bytes())
    assert rep["sleeves"] == 14 and rep["targets"] == 36 and rep["playbook"]
    assert db.scalar("SELECT COUNT(*) FROM platforms WHERE provider='bank' AND in_plan=0") == 3
    assert db.scalar("SELECT limit_gbp FROM platforms WHERE provider='freetrade'") == 78000
    assert playbook.evaluate()["signals"][6]["key"] == "nvda_dc"


# ============================================================================ corrections

def test_a_plan_correction_brings_holdings_to_what_the_broker_shows_once():
    """Freetrade's file leaves out transfers in and fund mergers. Allan's had 64 Procure
    Space shares that had become 25 Future of Defence, and ten holdings off in all."""
    pid = add_platform("Freetrade ISA")
    rows = [["Alphabet", "ORDER", "2025-03-03T10:00:00.000Z", "GBP", "1000", "BUY", "GOOGL", "US02079K3059",
             "100", "0", "10", "NASDAQ", "g1", "USD", "0"],
            ["Procure Space", "ORDER", "2023-05-17T10:00:00.000Z", "GBP", "249.9", "BUY", "UFOP", "IE00BLH3CV30",
             "3.9", "0", "64", "London Stock Exchange", "u1", "GBP", "0"]]
    freetrade.import_csv(pid, _ft(FT_NEW, rows))
    fix = {"id": "test-1", "platform": "Freetrade ISA", "holdings": {
        "GOOGL": {"shares": 34.09, "cost": 3357.44},
        "UFOP": {"shares": 0},
        "AMZN": {"shares": 22.37, "cost": 2574.35, "exchange": "US", "name": "Amazon"}}}
    rep = plan.import_plan({"whiskers_plan": 1, "corrections": fix})
    assert rep["corrections"]["missing"] == []
    def held():
        return {r["symbol"]: (round(r["quantity"], 4), round(r["cost"], 2)) for r in db.rows(
            "SELECT i.symbol, p.quantity, p.cost FROM positions p JOIN instruments i ON i.id=p.instrument_id")}
    want = {"GOOGL": (34.09, 3357.44), "AMZN": (22.37, 2574.35)}
    assert held() == want
    plan.import_plan({"whiskers_plan": 1, "corrections": fix})         # the same plan again
    freetrade.import_csv(pid, _ft(FT_NEW, rows))                        # and the same file again
    assert held() == want
    assert db.scalar("SELECT COUNT(*) FROM trades WHERE side='SET'") == 3


def test_a_correction_holds_even_when_the_file_sells_shares_it_never_saw_arrive():
    """Shares transferred in and later partly sold: the file's own count goes negative.
    Allan's RBTX came out at 53 instead of 351 before corrections replaced, not added."""
    pid = add_platform("Freetrade ISA")
    sell = [["Robotics", "ORDER", "2025-06-02T10:00:00.000Z", "GBP", "900", "SELL", "RBTX", "IE00BYZK4552",
             "9", "0", "100", "London Stock Exchange", "r1", "USD", "0"]]
    freetrade.import_csv(pid, _ft(FT_NEW, sell))
    plan.import_plan({"whiskers_plan": 1, "corrections": {"id": "t2", "platform": "Freetrade ISA",
                      "holdings": {"RBTX": {"shares": 351, "cost": 3226.27}}}})
    freetrade.import_csv(pid, _ft(FT_NEW, sell))
    assert db.one("SELECT quantity, cost FROM positions") == {"quantity": 351, "cost": 3226.27}


def test_one_company_spelt_two_ways_is_one_holding():
    """Freetrade writes BRK.B and Trading 212 BRK-B; without an ISIN to go on they were
    two rows on the Holdings screen."""
    a, b = add_platform("FT"), add_platform("T212", "trading212")
    assert add_position(a, "BRK.B") == add_position(b, "BRK-B")
    assert add_position(a, "SN.", "LSE") == add_position(b, "SN", "LSE")
    with db.tx() as c:
        x = symbols.upsert_instrument(c, "ABC", "US", None, "US0000000001", None)
        y = symbols.upsert_instrument(c, "ABC", "US", None, "US0000000002", None)
    assert x != y                                   # same ticker, different security


# ============================================================================ plans replacing a template

def test_a_plan_replaces_a_templates_leftover_sleeves_and_says_so():
    """Pressing "Start from a template" and then loading a real plan left the template's
    Global core, Industrials and Commodities & gold beside the plan's own."""
    sleeve("Global core", 40)
    old = sleeve("Industrials", 5)
    pid = add_platform()
    iid = add_position(pid, "XPO")
    db.execute("UPDATE instruments SET sleeve_id=? WHERE id=?", (old, iid))
    rep = plan.import_plan({"whiskers_plan": 1, "replace_sleeves": True,
                            "sleeves": [{"name": "Industrials & Materials", "target": 100}],
                            "instruments": {"XPO": {"sleeve": "Industrials & Materials"}}})
    assert sorted(rep["removed"]) == ["Global core", "Industrials"]
    assert (rep["sorted"], rep["unsorted"]) == (1, 0)
    assert [r["name"] for r in db.rows("SELECT name FROM sleeves")] == ["Industrials & Materials"]
    assert db.get_meta("plan_loaded")["at"]


# ============================================================================ the phone

def test_on_a_phone_the_api_answers_only_its_own_browser(server, monkeypatch):
    monkeypatch.setenv("WHISKERS_PHONE", "1")
    assert _req(server, "/api/bootstrap")[0] == 401                      # another app on the phone
    assert _req(server, "/api/bootstrap", headers={"X-Whiskers-Key": "0" * 48})[0] == 401
    key = config.phone_key()
    assert _req(server, "/api/bootstrap", headers={"X-Whiskers-Key": key})[0] == 200
    assert _req(server, "/")[0] == 200                                   # the page itself holds no data
    monkeypatch.delenv("WHISKERS_PHONE")
    assert _req(server, "/api/bootstrap")[0] == 200                      # the PC is unchanged


def test_the_app_can_be_installed_from_chrome(server):
    status, body, headers = _req(server, "/manifest.webmanifest")
    m = json.loads(body)
    assert headers["Content-Type"].startswith("application/manifest+json")
    assert m["display"] == "standalone" and {i["sizes"] for i in m["icons"]} == {"192x192", "512x512"}
    for i in m["icons"]:
        assert (STATIC / i["src"]).exists()


def test_the_phone_script_has_unix_line_endings():
    raw = (ROOT / "phone" / "start-whiskers.sh").read_bytes()
    assert b"\r\n" not in raw and raw.startswith(b"#!/data/data/com.termux")



def test_the_build_leaves_one_whiskers_exe_and_a_shortcut_to_it():
    """PyInstaller's scratch folder held a Whiskers.exe that cannot run by itself
    ("Failed to load Python DLL"), right beside the build script. That is the one Allan
    double-clicked. Scratch files now go to the temp folder, and the build makes a
    desktop shortcut to the real app."""
    bat = (ROOT / "build" / "build_windows.bat").read_text(encoding="utf-8")
    assert '--workpath "%TEMP%\\whiskers-build"' in bat and "--distpath dist" in bat
    assert "CreateShortcut" in bat and "dist\\Whiskers\\Whiskers.exe" in bat



def test_everything_a_plan_sets_can_be_saved_back_from_settings(server):
    """The plan gave "Being sold / not in plan" a band of 100; the Settings screen only
    took up to 50, so pressing Save the sleeves failed at that row and Allan's AI-share
    change on it was lost. Every sleeve is now saved back exactly as the screen sends it."""
    p = ROOT.parent / "out" / "Allan's plan.json"
    data = json.loads(p.read_text(encoding="utf-8")) if p.exists() else {
        "whiskers_plan": 1, "sleeves": [{"name": "Core", "target": 100},
                                        {"name": "Being sold / not in plan", "target": 0, "band": 100}]}
    plan.import_plan(data)
    ok = {"Content-Type": "application/json"}
    for s in db.rows("SELECT * FROM sleeves"):
        body = {"name": s["name"], "target": s["target"], "band": s["band"], "ai_share": 0,
                "is_cash": bool(s["is_cash"]), "colour": allocation.sleeves()[0]["colour"]}
        status, reply, _ = _req(server, f"/api/sleeves/{s['id']}", "PATCH", ok, json.dumps(body).encode())
        assert status == 200, (s["name"], reply)
    with pytest.raises(plan.PlanError):
        plan.import_plan({"whiskers_plan": 1, "sleeves": [{"name": "X", "target": 1, "band": 250}]})



# ============================================================================ cards and Monzo

def test_credit_cards_are_taken_off_cash_and_never_counted_as_protected(server):
    isa_ = add_platform("FT", cash=1000)
    add_position(isa_, "VWRP", "LSE", broker_value=9000, valued_at=db.now())
    ok = {"Content-Type": "application/json"}
    _req(server, "/api/platforms", "POST", ok, json.dumps({"name": "HSBC", "provider": "bank", "cash": 5000}).encode())
    status, body, _ = _req(server, "/api/platforms", "POST", ok,
                           json.dumps({"name": "Amex", "provider": "card", "owed": 1200}).encode())
    amex = json.loads(body)["id"]
    t = portfolio.totals()
    assert (t["total"], t["outside_cash"], t["owed"], t["outside_net"], t["everything"]) == \
        (10000, 5000, 1200, 3800, 13800)
    _req(server, f"/api/platforms/{amex}", "PATCH", ok, json.dumps({"owed": 300}).encode())
    assert portfolio.totals()["owed"] == 300
    card = next(p for p in portfolio.totals()["platforms"] if p["name"] == "Amex")
    assert card["kind"] == "card" and card["limit"] is None and not card["in_plan"]
    pos = portfolio.positions(); tot = portfolio.totals(pos)
    assert not [w for w in rules.evaluate(pos, tot, allocation.view(pos, tot)) if w["id"] == f"fscs:{amex}"]


def test_monzo_link_needs_a_state_this_copy_handed_out(server, monkeypatch):
    pid = add_platform("Monzo", "bank")
    with pytest.raises(monzo.MonzoError):
        monzo.begin(pid, "not-a-client", "x", server)
    out = monzo.begin(pid, "oauth2client_abc", "secret", server)
    q = dict(urllib.parse.parse_qsl(out["url"].split("?", 1)[1]))
    assert q["client_id"] == "oauth2client_abc" and q["response_type"] == "code"
    assert q["redirect_uri"] == f"http://127.0.0.1:{server}/oauth/monzo"
    # A forged link, even arriving as a cross-site navigation, links nothing.
    status, body, _ = _req(server, "/oauth/monzo?code=x&state=forged", headers={"Sec-Fetch-Site": "cross-site"})
    assert status == 200 and b"wasn" in body and not monzo.linked(pid)
    monkeypatch.setattr(monzo, "_post_token", lambda f: {"access_token": "a", "refresh_token": "r", "expires_in": 3600})
    status, body, _ = _req(server, f"/oauth/monzo?code=ok&state={q['state']}", headers={"Sec-Fetch-Site": "cross-site"})
    assert b"is linked" in body and monzo.linked(pid)
    status, body, _ = _req(server, f"/oauth/monzo?code=ok&state={q['state']}")   # a state works once
    assert b"wasn" in body


def test_monzo_balance_counts_the_account_and_its_pots(monkeypatch):
    pid = add_platform("Monzo", "bank")
    import time as _t
    from folio import net, security
    security.put(f"monzo:{pid}", {"client_id": "oauth2client_x", "client_secret": "s",
                                  "access": "tok", "refresh": "r", "expires": _t.time() + 3600})
    replies = {"/accounts": {"accounts": [{"id": "acc_1", "type": "uk_retail", "closed": False},
                                          {"id": "acc_2", "type": "uk_retail_joint", "closed": False}]},
               "/balance?account_id=acc_1": {"balance": 12345, "total_balance": 512345, "currency": "GBP"}}
    monkeypatch.setattr(net, "get", lambda url, **k: json.dumps(replies[url.split("monzo.com", 1)[1]]).encode())
    got = monzo.sync(pid)
    assert got["balance"] == 5123.45 and got["accounts"] == 1          # the joint account is left out
    assert db.scalar("SELECT cash FROM platforms WHERE id=?", (pid,)) == 5123.45

    def refuse(url, **k):
        raise net.FetchError("forbidden.verification_required", status=403)
    monkeypatch.setattr(net, "get", refuse)
    with pytest.raises(monzo.MonzoError, match="approve"):
        monzo.sync(pid)


def test_monzo_client_can_only_read():
    src = (ROOT / "folio" / "brokers" / "monzo.py").read_text(encoding="utf-8")
    assert src.count('method="POST"') == 1 and "/oauth2/token" in src
    for verb in ("/pots/", "/deposit", "/withdraw", "/feed", "PUT", "PATCH", "DELETE"):
        assert verb not in src.replace("DELETE FROM", "")



# ============================================================================ trends, not jolts

def test_one_bad_day_does_not_turn_a_light():
    """Lights are judged on a 5-day average: a one-day spike stays green, a week of it
    turns the light, and the light then says since when and which way it is heading."""
    vix = next(d for d in signals.DEFS if d["key"] == "vix")
    assert _levels(vix, [15.0] * 40 + [35.0])[-1][1] == "green"
    assert _levels(vix, [15.0] * 40 + [35.0] * 5)[-1][1] == "red"
    lv = _levels(vix, [15.0] * 60 + [22.0] * 30)
    t = signals._trend(vix, lv)
    assert t["direction"] == "worsening" and lv[-1][1] == "amber"


def test_the_yield_curve_warns_most_when_it_turns_back_up():
    curve = next(d for d in signals.DEFS if d["key"] == "curve")
    assert _levels(curve, [0.8] * 300)[-1][1] == "green"
    assert _levels(curve, [0.8] * 100 + [-0.5] * 200)[-1][1] == "amber"
    assert _levels(curve, [0.8] * 50 + [-0.5] * 200 + [0.3] * 10)[-1][1] == "red"


def test_the_big_picture_hides_amounts_and_never_forecasts():
    market = {"lights": [{"key": "sox_fall", "title": "Chip stocks", "group": "ai", "level": "amber",
                          "since": "2026-09-01", "direction": "worsening"}],
              "groups": {"ai": {"label": "Watch"}, "market": {"label": "Calm"}}}
    exp = {"ai_pct": 0.3, "ai_value": 42685.0, "top": [{"name": "Invesco S&P 500"}]}
    b = narrative.big_picture(market, None, exp, {"sleeves": []}, {"platforms": []})
    assert "[[£42,685]]" in b["portfolio"] and "predicts" in b["portfolio"]
    assert "Chip stocks: amber since 1 Sep, still getting worse" in b["ai"]



# ============================================================================ the plan builder

@pytest.mark.parametrize("answers", [
    {}, {"horizon": "under5", "fall": "sell", "emergency": "no"},
    {"horizon": "over20", "fall": "buy", "db": "yes", "style": "hands",
     "themes": list(planbuilder.THEMES), "ai_cap": "30", "platforms": ["freetrade", "trading212"],
     "buffer": "78000"},
    {"style": "simple", "themes": ["tech"]},
])
def test_every_built_plan_adds_up_and_loads(answers):
    out = planbuilder.build(answers)
    sleeves = out["plan"]["sleeves"]
    assert abs(sum(s["target"] for s in sleeves) - 100) < 1e-9
    assert all(0 < s["target"] <= 100 for s in sleeves) and sum(1 for s in sleeves if s.get("is_cash")) == 1
    assert out["why"][-1].startswith("This is a starting point")
    sleeve("Global core", 40)                                     # a template left behind
    rep = plan.import_plan(out["plan"])
    assert rep["removed"] == ["Global core"] and allocation.view([], {"total": 0})["targets_add_up"]


def test_the_builder_keeps_more_safe_when_the_money_is_needed_soon():
    soon = planbuilder.build({"horizon": "under5", "fall": "worry"})
    later = planbuilder.build({"horizon": "over20", "fall": "hold", "db": "yes"})
    assert soon["safe"] > 50 > later["safe"]
    assert later["plan"]["rules"] == [] and not later["plan"]["platforms"]



def test_a_retiree_drawing_income_gets_buckets_and_inflation_protection():
    out = planbuilder.build({"horizon": "now", "withdraw": "4", "age": "65plus", "fall": "hold",
                             "experience": "long", "prefer": "balanced"})
    got = {s["name"]: s["target"] for s in out["plan"]["sleeves"]}
    assert got["Cash"] == 8                          # two years of 4% withdrawals
    assert got["Index-linked gilts"] > 0 and got["Gilts and bonds"] > 0
    assert abs(sum(got.values()) - 100) < 1e-9
    assert "Spend from the Cash sleeve" in out["plan"]["plans"]["crash"]
    high = planbuilder.build({"horizon": "now", "withdraw": "5"})
    assert any("3.9%" in w for w in high["why"])


def test_a_secure_pension_never_pushes_past_what_someone_is_willing_to_take():
    nervous = planbuilder.build({"horizon": "over20", "guaranteed": "most", "fall": "sell",
                                 "experience": "new", "prefer": "steady"})
    assert nervous["growth"] <= 30 + 1e-9          # willingness caps it, however secure they are
    bold = planbuilder.build({"horizon": "over20", "guaranteed": "most", "fall": "buy",
                              "experience": "long", "prefer": "growth"})
    assert bold["growth"] > 90 and all(len(x) == 2 for x in bold["sources"])



# ============================================================================ the spreadsheet export

def test_the_export_has_every_sheet_with_real_figures_in_them():
    sleeve("Technology & AI", 60, ai_share=100)
    sleeve("Cash", 40, is_cash=True)
    pid = add_platform("Freetrade ISA", cash=500)
    iid = add_position(pid, "NVDA", qty=10, cost=900, broker_value=1200, valued_at=db.now(), name="NVIDIA")
    db.execute("UPDATE instruments SET sleeve_id=(SELECT id FROM sleeves WHERE name='Technology & AI') "
              "WHERE id=?", (iid,))
    with db.tx() as c:
        c.execute("INSERT INTO dividends(platform_id,instrument_id,paid_on,amount_gbp,fp) VALUES"
                  "(?,?,'2026-06-01',12.5,'d1')", (pid, iid))
        c.execute("INSERT INTO trades(platform_id,instrument_id,traded_on,side,quantity,value_gbp,fp) "
                  "VALUES(?,?,'2026-01-01','BUY',10,900,'t1')", (pid, iid))
        c.execute("INSERT INTO cash_moves(platform_id,happened_on,kind,amount_gbp,fp) VALUES"
                  "(?,'2026-01-01','DEPOSIT',1000,'c1')", (pid,))

    import io
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(export.build()))
    assert wb.sheetnames == ["Summary", "Holdings", "Allocation", "Dividends", "Trades", "Cash movements"]

    hold = {c.value: i for i, c in enumerate(next(wb["Holdings"].iter_rows(min_row=1, max_row=1)), 1)}
    row = next(wb["Holdings"].iter_rows(min_row=2, max_row=2))
    assert row[hold["Name"] - 1].value == "NVIDIA"
    assert row[hold["Cost (£)"] - 1].value == 900 and row[hold["Value (£)"] - 1].value == 1200
    assert row[hold["Shares"] - 1].value == 10
    assert row[hold["AI share"] - 1].value == pytest.approx(1.0)       # a fraction, not 100
    assert wb["Holdings"]["A1"].fill.fgColor.rgb == "000B6E6E"          # the brand colour header

    alloc = {c.value: i for i, c in enumerate(next(wb["Allocation"].iter_rows(min_row=1, max_row=1)), 1)}
    tech_row = next(r for r in wb["Allocation"].iter_rows(min_row=2) if r[alloc["Sleeve"] - 1].value == "Technology & AI")
    assert tech_row[alloc["Target"] - 1].value == pytest.approx(0.6)
    assert tech_row[alloc["Target"] - 1].number_format == "0.0%"

    div_row = next(wb["Dividends"].iter_rows(min_row=2, max_row=2))
    assert div_row[2].value == "NVIDIA" and div_row[4].value == 12.5
    assert next(wb["Trades"].iter_rows(min_row=2, max_row=2))[4].value == "Bought"
    assert next(wb["Cash movements"].iter_rows(min_row=2, max_row=2))[2].value == "Paid in"


def test_an_empty_copy_exports_without_error():
    import io
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(export.build()))
    assert wb["Holdings"].max_row == 1 and wb["Summary"]["A1"].value.startswith("Whiskers")


def test_the_export_route_sends_a_real_xlsx_as_an_attachment(server):
    status, body, headers = _req(server, "/api/export/xlsx")
    assert status == 200
    assert headers["Content-Type"] == \
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert "Whiskers export" in headers["Content-Disposition"] and ".xlsx" in headers["Content-Disposition"]
    assert body[:2] == b"PK"                       # every xlsx is a zip file
    import io
    from openpyxl import load_workbook
    load_workbook(io.BytesIO(body))                # raises if the bytes aren't a real workbook


def test_a_treasury_bill_shows_its_value_source_in_plain_words():
    sleeve("Cash", 100, is_cash=True)
    pid = add_platform("Freetrade ISA")
    rows = [_bill_row("GB00BSGQD757", (dt.date.today() - dt.timedelta(days=5)).isoformat(), 1010)]
    freetrade.import_csv(pid, _ft(FT_NEW, rows))
    import io
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(export.build()))
    hold = {c.value: i - 1 for i, c in enumerate(next(wb["Holdings"].iter_rows(min_row=1, max_row=1)), 1)}
    row = next(wb["Holdings"].iter_rows(min_row=2, max_row=2))
    assert row[hold["Value source"]].value == "Treasury bill, at cost"
